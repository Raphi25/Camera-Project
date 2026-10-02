#!/usr/bin/env python3
"""Analyze repeated low-light images grouped in folders such as ``35LUX``.

Metrics are calculated over each complete image by default.  Noise is
estimated as the standard deviation of the high-frequency residual after a
small Gaussian blur; SNR is 20*log10(mean luminance / noise sigma).  Sharpness
is the variance of the luminance Laplacian.
"""

from __future__ import annotations

import argparse
import math
import re
from pathlib import Path

import cv2
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


DEFAULT_INPUT = Path(r"C:\Users\LAB-ADMIN\Pictures\SD_Card_Media\SC2336_L")
IMAGE_EXTENSIONS = {".jpg", ".jpeg", ".png", ".bmp", ".tif", ".tiff"}


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("input", nargs="?", type=Path, default=DEFAULT_INPUT)
    parser.add_argument("--output", type=Path, default=Path("analysis_output"))
    parser.add_argument("--crop", nargs=4, type=float, metavar=("X", "Y", "W", "H"),
                        default=(0.0, 0.0, 1.0, 1.0),
                        help="Optional normalized crop coordinates (default: complete image)")
    parser.add_argument("--noise-blur-sigma", type=float, default=1.0)
    return parser.parse_args()


def lux_from_folder(folder: Path) -> float | None:
    match = re.search(r"([0-9]+(?:\.[0-9]+)?)\s*LUX", folder.name, re.IGNORECASE)
    return float(match.group(1)) if match else None


def crop_bounds(shape: tuple[int, ...], crop: tuple[float, float, float, float]) -> tuple[int, int, int, int]:
    height, width = shape[:2]
    x, y, w, h = crop
    if not (0 <= x < 1 and 0 <= y < 1 and 0 < w <= 1 and 0 < h <= 1 and x + w <= 1 and y + h <= 1):
        raise ValueError("Crop values must describe a normalized rectangle inside the image")
    x0, y0 = round(x * width), round(y * height)
    x1, y1 = round((x + w) * width), round((y + h) * height)
    if x1 <= x0 or y1 <= y0:
        raise ValueError("Crop is empty at this image resolution")
    return x0, y0, x1, y1


def luminance(bgr: np.ndarray) -> np.ndarray:
    return cv2.cvtColor(bgr, cv2.COLOR_BGR2GRAY).astype(np.float32)


def image_metrics(gray: np.ndarray, blur_sigma: float) -> dict[str, float]:
    mean = float(np.mean(gray))
    smooth = cv2.GaussianBlur(gray, (0, 0), blur_sigma)
    noise_std = float(np.std(gray - smooth, ddof=1))
    snr_db = float(20 * math.log10(mean / noise_std)) if mean > 0 and noise_std > 0 else math.nan
    sharpness = float(cv2.Laplacian(gray, cv2.CV_32F).var())
    return {"mean_brightness": mean, "noise_std": noise_std, "snr_db": snr_db,
            "sharpness_laplacian_variance": sharpness}


def style_workbook(path: Path) -> None:
    wb = load_workbook(path)
    navy, blue, pale = "17365D", "5B9BD5", "D9EAF7"
    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        ws.sheet_view.showGridLines = False
        for cell in ws[1]:
            cell.fill = PatternFill("solid", fgColor=navy)
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(horizontal="center")
        for col in range(1, ws.max_column + 1):
            values = [str(ws.cell(r, col).value or "") for r in range(1, min(ws.max_row, 100) + 1)]
            ws.column_dimensions[get_column_letter(col)].width = min(42, max(11, max(map(len, values), default=0) + 2))
        for row in range(2, ws.max_row + 1):
            if row % 2 == 0:
                for cell in ws[row]:
                    cell.fill = PatternFill("solid", fgColor=pale)
    ws = wb["Lux Summary"]
    headers = {c.value: c.column for c in ws[1]}
    for metric, title, anchor in [
        ("mean_brightness_mean", "Mean brightness vs lux", "N2"),
        ("noise_std_mean", "Estimated noise vs lux", "N18"),
        ("snr_db_mean", "SNR vs lux (dB)", "N34"),
        ("sharpness_laplacian_variance_mean", "Sharpness vs lux", "N50"),
    ]:
        chart = LineChart()
        chart.title, chart.style, chart.height, chart.width = title, 13, 7.2, 13
        chart.x_axis.title, chart.y_axis.title = "Illuminance (lux)", title.split(" vs lux")[0]
        data = Reference(ws, min_col=headers[metric], min_row=1, max_row=ws.max_row)
        cats = Reference(ws, min_col=headers["lux"], min_row=2, max_row=ws.max_row)
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(cats)
        chart.legend = None
        chart.series[0].graphicalProperties.line.solidFill = blue
        chart.series[0].graphicalProperties.line.width = 24000
        ws.add_chart(chart, anchor)
    wb.save(path)


def plot_summary(summary: pd.DataFrame, output: Path) -> None:
    # Every recognized lux folder becomes a plotted point. A logarithmic axis
    # provides useful spacing while explicit ticks show every measured value.
    plotted = summary.copy()
    measured_ticks = sorted(plotted["lux"].astype(float).unique())
    metrics = [
        ("mean_brightness_mean", "Mean brightness", "8-bit luminance"),
        ("noise_std_mean", "Estimated noise", "High-pass residual σ"),
        ("snr_db_mean", "Signal-to-noise ratio", "SNR (dB)"),
        ("sharpness_laplacian_variance_mean", "Sharpness", "Laplacian variance"),
        ("repeat_pixel_std_mean", "Repeat variability", "Mean pixelwise σ"),
    ]
    fig, axes = plt.subplots(3, 2, figsize=(12, 12), constrained_layout=True)
    for ax, (column, title, ylabel) in zip(axes.flat, metrics):
        error_col = column.replace("_mean", "_std")
        errors = plotted[error_col] if error_col in plotted else None
        ax.errorbar(plotted["lux"], plotted[column], yerr=errors, marker="o", capsize=4, linewidth=2)
        ax.set_xscale("log")
        ax.set_xticks(measured_ticks)
        ax.set_xticklabels([f"{lux:g}" for lux in measured_ticks])
        # Present the test sequence from brightest to darkest illumination.
        ax.invert_xaxis()
        ax.tick_params(axis="x", labelrotation=45)
        ax.set_title(title)
        ax.set_xlabel("Illuminance (lux, log scale)")
        ax.set_ylabel(ylabel)
        ax.grid(True, alpha=0.25)
    axes.flat[-1].axis("off")
    fig.suptitle("Image quality versus illuminance", fontsize=16)
    fig.savefig(output / "image_quality_vs_lux.png", dpi=180)
    plt.close(fig)


def main() -> None:
    args = parse_args()
    input_root, output = args.input.resolve(), args.output.resolve()
    if not input_root.is_dir():
        raise SystemExit(f"Input folder not found: {input_root}")
    output.mkdir(parents=True, exist_ok=True)
    averages_dir, variability_dir = output / "average_images", output / "repeat_std_maps"
    averages_dir.mkdir(exist_ok=True)
    variability_dir.mkdir(exist_ok=True)

    groups = []
    for folder in input_root.iterdir():
        lux = lux_from_folder(folder) if folder.is_dir() else None
        if lux is not None:
            files = sorted(p for p in folder.iterdir() if p.is_file() and p.suffix.lower() in IMAGE_EXTENSIONS)
            if files:
                groups.append((lux, folder.name, files))
    if not groups:
        raise SystemExit("No folders named like '35LUX' containing images were found")

    rows, summary_rows = [], []
    reference_shape, reference_bounds = None, None
    for lux, group_name, files in sorted(groups):
        print(f"Processing {group_name}: {len(files)} image(s)", flush=True)
        color_crops, gray_crops, group_metrics = [], [], []
        for repeat, path in enumerate(files, 1):
            image = cv2.imread(str(path), cv2.IMREAD_COLOR)
            if image is None:
                raise SystemExit(f"Could not decode image: {path}")
            if reference_shape is None:
                reference_shape = image.shape
                reference_bounds = crop_bounds(image.shape, tuple(args.crop))
            if image.shape != reference_shape:
                raise SystemExit(f"Image dimensions differ: {path} is {image.shape}, expected {reference_shape}")
            x0, y0, x1, y1 = reference_bounds
            color_crop = image[y0:y1, x0:x1]
            gray_crop = luminance(color_crop)
            metrics = image_metrics(gray_crop, args.noise_blur_sigma)
            rows.append({"lux": lux, "group": group_name, "repeat": repeat, "filename": path.name,
                         "width_px": image.shape[1], "height_px": image.shape[0],
                         "crop_x": x0, "crop_y": y0, "crop_width": x1 - x0, "crop_height": y1 - y0,
                         **metrics})
            color_crops.append(color_crop.astype(np.float32))
            gray_crops.append(gray_crop)
            group_metrics.append(metrics)

        color_stack, gray_stack = np.stack(color_crops), np.stack(gray_crops)
        average_color = np.clip(np.mean(color_stack, axis=0), 0, 255).astype(np.uint8)
        # Sample standard deviation requires at least two repeat images.
        has_repeats = len(files) > 1
        pixel_std = np.std(gray_stack, axis=0, ddof=1) if has_repeats else np.zeros_like(gray_stack[0])
        cv2.imwrite(str(averages_dir / f"{group_name}_average.png"), average_color)
        std_display = np.clip(pixel_std / max(float(pixel_std.max()), 1e-9) * 255, 0, 255).astype(np.uint8)
        cv2.imwrite(str(variability_dir / f"{group_name}_repeat_std.png"), std_display)
        item = {"lux": lux, "group": group_name, "image_count": len(files),
                "repeat_pixel_std_mean": float(pixel_std.mean()) if has_repeats else math.nan,
                "repeat_pixel_std_median": float(np.median(pixel_std)) if has_repeats else math.nan,
                "repeat_pixel_std_max": float(pixel_std.max()) if has_repeats else math.nan}
        for key in group_metrics[0]:
            values = np.array([m[key] for m in group_metrics], dtype=float)
            item[f"{key}_mean"] = float(np.mean(values))
            item[f"{key}_std"] = float(np.std(values, ddof=1)) if has_repeats else math.nan
        summary_rows.append(item)

    details = pd.DataFrame(rows).sort_values(["lux", "repeat"])
    summary = pd.DataFrame(summary_rows).sort_values("lux")
    details.to_csv(output / "per_image_metrics.csv", index=False, float_format="%.6f")
    summary.to_csv(output / "lux_summary.csv", index=False, float_format="%.6f")
    metadata = pd.DataFrame([
        ["Input folder", str(input_root)], ["Crop (normalized x,y,w,h)", ", ".join(map(str, args.crop))],
        ["Crop (pixels x,y,w,h)", f"{reference_bounds[0]}, {reference_bounds[1]}, {reference_bounds[2]-reference_bounds[0]}, {reference_bounds[3]-reference_bounds[1]}"],
        ["Noise definition", f"Std. dev. of grayscale residual after Gaussian blur (sigma={args.noise_blur_sigma})"],
        ["SNR definition", "20*log10(mean grayscale brightness / estimated noise std.)"],
        ["Sharpness definition", "Variance of grayscale Laplacian"],
        ["Repeat variability", "Pixelwise grayscale sample std. dev. across all images in each lux folder (requires at least two)"],
    ], columns=["Setting", "Value"])
    excel_path = output / "low_light_analysis.xlsx"
    with pd.ExcelWriter(excel_path, engine="openpyxl") as writer:
        summary.to_excel(writer, sheet_name="Lux Summary", index=False)
        details.to_excel(writer, sheet_name="Per Image", index=False)
        metadata.to_excel(writer, sheet_name="Method", index=False)
    style_workbook(excel_path)
    plot_summary(summary, output)
    print(f"Analyzed {len(details)} images across {len(summary)} lux levels")
    print(f"Crop pixels (x, y, width, height): {reference_bounds[0]}, {reference_bounds[1]}, {reference_bounds[2]-reference_bounds[0]}, {reference_bounds[3]-reference_bounds[1]}")
    print(f"Results: {output}")


if __name__ == "__main__":
    main()
