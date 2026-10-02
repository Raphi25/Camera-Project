#!/usr/bin/env python3
"""Analyze separate grayscale and colour-chart captures across lux levels.

Expected layout::

    dataset/grayscale/35LUX/*.jpg
    dataset/colorchecker/35LUX/*.jpg

Chart corners and grid geometry are stored in a JSON configuration file.  This
is an ISO-inspired laboratory aid, not a declaration of ISO 19093 compliance.
"""

from __future__ import annotations

import argparse
import json
import math
import re
from pathlib import Path

import cv2
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


IMAGE_EXTENSIONS = {".jpg", ".jpeg", ".png", ".bmp", ".tif", ".tiff"}
DEFAULT_CONFIG = {
    "grayscale": {
        "corners": [[0.05, 0.15], [0.95, 0.15], [0.95, 0.85], [0.05, 0.85]],
        "rows": 1,
        "columns": 16,
        "reverse_patch_order": False,
    },
    "colorchecker": {
        "corners": [[0.05, 0.05], [0.95, 0.05], [0.95, 0.95], [0.05, 0.95]],
        "rows": 6,
        "columns": 4,
        "reverse_patch_order": False,
    },
    "patch_inset_fraction": 0.20,
    "align_repeats": True,
}


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("input", type=Path, help="Root containing grayscale/ and colorchecker/")
    parser.add_argument("--config", type=Path, default=Path("chart_config.json"))
    parser.add_argument("--output", type=Path, default=Path("chart_analysis_output"))
    parser.add_argument("--gray-reference", type=Path,
                        help="CSV: patch_id and relative_luminance (or reflectance)")
    parser.add_argument("--color-reference", type=Path,
                        help="CSV: patch_id, L_ref, a_ref, b_ref; optional patch_name, neutral")
    parser.add_argument("--write-default-config", action="store_true",
                        help="Write a documented default chart_config.json and exit")
    return parser.parse_args()


def write_config(path: Path) -> None:
    path.write_text(json.dumps(DEFAULT_CONFIG, indent=2), encoding="utf-8")
    print(f"Wrote {path.resolve()}")


def load_config(path: Path) -> dict:
    if not path.exists():
        write_config(path)
        raise SystemExit("Edit the generated chart_config.json corners, then run again.")
    config = json.loads(path.read_text(encoding="utf-8"))
    for chart in ("grayscale", "colorchecker"):
        if chart not in config or len(config[chart].get("corners", [])) != 4:
            raise SystemExit(f"Config must contain four normalized corners for {chart}")
    inset = float(config.get("patch_inset_fraction", 0.20))
    if not 0 <= inset < 0.45:
        raise SystemExit("patch_inset_fraction must be between 0 and 0.45")
    return config


def lux_from_folder(folder: Path) -> float | None:
    match = re.search(r"([0-9]+(?:\.[0-9]+)?)\s*LUX", folder.name, re.IGNORECASE)
    return float(match.group(1)) if match else None


def collect_groups(root: Path) -> list[tuple[float, str, list[Path]]]:
    groups = []
    if not root.is_dir():
        return groups
    for folder in root.iterdir():
        lux = lux_from_folder(folder) if folder.is_dir() else None
        if lux is None:
            continue
        files = sorted(p for p in folder.iterdir() if p.is_file() and p.suffix.lower() in IMAGE_EXTENSIONS)
        if files:
            groups.append((lux, folder.name, files))
    return sorted(groups)


def warp_chart(image: np.ndarray, spec: dict) -> tuple[np.ndarray, np.ndarray]:
    height, width = image.shape[:2]
    points = np.array([[x * width, y * height] for x, y in spec["corners"]], dtype=np.float32)
    tl, tr, br, bl = points
    out_w = max(32, round(max(np.linalg.norm(tr - tl), np.linalg.norm(br - bl))))
    out_h = max(32, round(max(np.linalg.norm(bl - tl), np.linalg.norm(br - tr))))
    destination = np.array([[0, 0], [out_w - 1, 0], [out_w - 1, out_h - 1], [0, out_h - 1]], np.float32)
    matrix = cv2.getPerspectiveTransform(points, destination)
    return cv2.warpPerspective(image, matrix, (out_w, out_h)), points


def align_translation(reference: np.ndarray, moving: np.ndarray) -> tuple[np.ndarray, float, float, float]:
    ref = cv2.cvtColor(reference, cv2.COLOR_BGR2GRAY).astype(np.float32) / 255.0
    mov = cv2.cvtColor(moving, cv2.COLOR_BGR2GRAY).astype(np.float32) / 255.0
    warp = np.eye(2, 3, dtype=np.float32)
    try:
        score, warp = cv2.findTransformECC(ref, mov, warp, cv2.MOTION_TRANSLATION,
                                           (cv2.TERM_CRITERIA_EPS | cv2.TERM_CRITERIA_COUNT, 100, 1e-6))
        aligned = cv2.warpAffine(moving, warp, (reference.shape[1], reference.shape[0]),
                                 flags=cv2.INTER_LINEAR | cv2.WARP_INVERSE_MAP,
                                 borderMode=cv2.BORDER_REFLECT)
        return aligned, float(warp[0, 2]), float(warp[1, 2]), float(score)
    except cv2.error:
        return moving, 0.0, 0.0, math.nan


def patch_regions(image: np.ndarray, rows: int, columns: int, inset: float, reverse: bool):
    height, width = image.shape[:2]
    patch_id = 0
    regions = []
    for row in range(rows):
        for col in range(columns):
            patch_id += 1
            x0, x1 = round(col * width / columns), round((col + 1) * width / columns)
            y0, y1 = round(row * height / rows), round((row + 1) * height / rows)
            dx, dy = round((x1 - x0) * inset), round((y1 - y0) * inset)
            sample = image[y0 + dy:y1 - dy, x0 + dx:x1 - dx]
            output_id = rows * columns + 1 - patch_id if reverse else patch_id
            regions.append((output_id, row + 1, col + 1, x0 + dx, y0 + dy, x1 - dx, y1 - dy, sample))
    return sorted(regions, key=lambda item: item[0])


def srgb_bgr_to_lab(mean_bgr: np.ndarray) -> tuple[float, float, float]:
    pixel = np.clip(mean_bgr, 0, 255).astype(np.uint8).reshape(1, 1, 3)
    lab = cv2.cvtColor(pixel, cv2.COLOR_BGR2LAB)[0, 0].astype(float)
    return lab[0] * 100.0 / 255.0, lab[1] - 128.0, lab[2] - 128.0


def measure_patch(sample: np.ndarray) -> dict[str, float]:
    pixels = sample.reshape(-1, 3).astype(np.float32)
    mean_bgr = pixels.mean(axis=0)
    std_bgr = pixels.std(axis=0, ddof=1)
    gray = cv2.cvtColor(sample, cv2.COLOR_BGR2GRAY).astype(np.float32)
    mean_y, spatial_std = float(gray.mean()), float(gray.std(ddof=1))
    residual = gray - cv2.GaussianBlur(gray, (0, 0), 1.0)
    highpass_std = float(residual.std(ddof=1))
    snr_db = 20 * math.log10(mean_y / spatial_std) if mean_y > 0 and spatial_std > 0 else math.nan
    l_star, a_star, b_star = srgb_bgr_to_lab(mean_bgr)
    return {
        "mean_R": float(mean_bgr[2]), "mean_G": float(mean_bgr[1]), "mean_B": float(mean_bgr[0]),
        "std_R": float(std_bgr[2]), "std_G": float(std_bgr[1]), "std_B": float(std_bgr[0]),
        "mean_luminance": mean_y, "spatial_noise_std": spatial_std,
        "highpass_noise_std": highpass_std, "snr_db": snr_db,
        "L_star": l_star, "a_star": a_star, "b_star": b_star,
        "chroma": float(math.hypot(a_star, b_star)),
        "black_clip_percent": float(100 * np.mean(gray <= 1)),
        "white_clip_percent": float(100 * np.mean(gray >= 254)),
    }


def annotate(image: np.ndarray, regions, output: Path) -> None:
    canvas = image.copy()
    for patch_id, _, _, x0, y0, x1, y1, _ in regions:
        cv2.rectangle(canvas, (x0, y0), (x1 - 1, y1 - 1), (0, 255, 0), 1)
        cv2.putText(canvas, str(patch_id), (x0 + 3, y0 + 16), cv2.FONT_HERSHEY_SIMPLEX,
                    0.45, (0, 0, 255), 1, cv2.LINE_AA)
    cv2.imwrite(str(output), canvas)


def analyze_chart(chart_name: str, groups, spec: dict, inset: float, align: bool, diagnostic_dir: Path):
    rows = []
    for lux, group, files in groups:
        reference = None
        per_repeat_patch_means = []
        for repeat, path in enumerate(files, 1):
            original = cv2.imread(str(path), cv2.IMREAD_COLOR)
            if original is None:
                raise SystemExit(f"Could not decode {path}")
            warped, source_points = warp_chart(original, spec)
            shift_x = shift_y = 0.0
            alignment_score = math.nan
            if reference is None:
                reference = warped
            elif align:
                warped, shift_x, shift_y, alignment_score = align_translation(reference, warped)
            regions = patch_regions(warped, int(spec["rows"]), int(spec["columns"]), inset,
                                    bool(spec.get("reverse_patch_order", False)))
            if repeat == 1:
                annotate(warped, regions, diagnostic_dir / f"{chart_name}_{group}_patches.jpg")
            repeat_values = []
            for patch_id, row, col, x0, y0, x1, y1, sample in regions:
                metrics = measure_patch(sample)
                repeat_values.append(metrics["mean_luminance"])
                rows.append({
                    "chart": chart_name, "lux": lux, "group": group, "repeat": repeat,
                    "filename": path.name, "patch_id": patch_id, "patch_row": row, "patch_column": col,
                    "roi_x": x0, "roi_y": y0, "roi_width": x1 - x0, "roi_height": y1 - y0,
                    "patch_width_percent_image_height": 100 * (x1 - x0) / original.shape[0],
                    "patch_height_percent_image_height": 100 * (y1 - y0) / original.shape[0],
                    "alignment_x_px": shift_x, "alignment_y_px": shift_y,
                    "alignment_score": alignment_score, **metrics,
                })
            per_repeat_patch_means.append(repeat_values)
        repeat_array = np.asarray(per_repeat_patch_means, dtype=float)
        repeat_std = np.std(repeat_array, axis=0, ddof=1)
        for record in rows:
            if record["chart"] == chart_name and record["lux"] == lux:
                record["repeat_mean_luminance_std"] = float(repeat_std[int(record["patch_id"]) - 1])
    return pd.DataFrame(rows)


def load_reference(path: Path | None, required: list[str]) -> pd.DataFrame | None:
    if path is None:
        return None
    table = pd.read_csv(path)
    missing = [name for name in required if name not in table.columns]
    if missing:
        raise SystemExit(f"{path} is missing columns: {', '.join(missing)}")
    if table["patch_id"].duplicated().any():
        raise SystemExit(f"{path} contains duplicate patch_id values")
    return table


def summarize_gray(details: pd.DataFrame, reference: pd.DataFrame | None) -> pd.DataFrame:
    numeric = ["mean_R", "mean_G", "mean_B", "mean_luminance", "spatial_noise_std",
               "highpass_noise_std", "snr_db", "black_clip_percent", "white_clip_percent"]
    summary = details.groupby(["lux", "patch_id"], as_index=False)[numeric].agg(["mean", "std"])
    summary.columns = ["_".join(filter(None, map(str, col))) for col in summary.columns]
    summary = summary.rename(columns={"lux_": "lux", "patch_id_": "patch_id"})
    if reference is not None:
        ref = reference.copy()
        if "relative_luminance" not in ref and "reflectance" in ref:
            ref["relative_luminance"] = ref["reflectance"]
        if "relative_luminance" not in ref:
            raise SystemExit("Gray reference needs relative_luminance or reflectance")
        summary = summary.merge(ref, on="patch_id", how="left", validate="many_to_one")
        summary["scene_luminance_proxy"] = summary["lux"] * summary["relative_luminance"]
        summary["log10_scene_luminance_proxy"] = np.log10(summary["scene_luminance_proxy"].where(lambda x: x > 0))
    return summary


def summarize_color(details: pd.DataFrame, reference: pd.DataFrame | None) -> pd.DataFrame:
    numeric = ["mean_R", "mean_G", "mean_B", "mean_luminance", "spatial_noise_std", "snr_db",
               "L_star", "a_star", "b_star", "chroma", "black_clip_percent", "white_clip_percent"]
    summary = details.groupby(["lux", "patch_id"], as_index=False)[numeric].agg(["mean", "std"])
    summary.columns = ["_".join(filter(None, map(str, col))) for col in summary.columns]
    summary = summary.rename(columns={"lux_": "lux", "patch_id_": "patch_id"})
    max_lux = float(summary["lux"].max())
    baseline = summary.loc[summary["lux"] == max_lux, ["patch_id", "chroma_mean"]].rename(
        columns={"chroma_mean": "chroma_at_max_lux"})
    summary = summary.merge(baseline, on="patch_id", how="left")
    summary["chroma_retention_vs_max_lux"] = summary["chroma_mean"] / summary["chroma_at_max_lux"].replace(0, np.nan)
    if reference is not None:
        summary = summary.merge(reference, on="patch_id", how="left", validate="many_to_one")
        summary["delta_E76"] = np.sqrt((summary["L_star_mean"] - summary["L_ref"]) ** 2 +
                                             (summary["a_star_mean"] - summary["a_ref"]) ** 2 +
                                             (summary["b_star_mean"] - summary["b_ref"]) ** 2)
    return summary


def plot_gray(summary: pd.DataFrame, output: Path) -> None:
    x_column = "log10_scene_luminance_proxy" if "log10_scene_luminance_proxy" in summary else "patch_id"
    x_label = "log10(lux × relative luminance) proxy" if x_column.startswith("log10") else "Gray patch ID"
    fig, axes = plt.subplots(2, 2, figsize=(12, 9), constrained_layout=True)
    for lux, group in summary.groupby("lux"):
        group = group.sort_values(x_column)
        axes[0, 0].plot(group[x_column], group["mean_luminance_mean"], marker="o", label=f"{lux:g} lux")
        axes[0, 1].plot(group["mean_luminance_mean"], group["spatial_noise_std_mean"], marker="o")
        axes[1, 0].plot(group["mean_luminance_mean"], group["snr_db_mean"], marker="o")
        axes[1, 1].plot(group["patch_id"], group["mean_luminance_std"], marker="o")
    axes[0, 0].set(title="OECF / tonal response", xlabel=x_label, ylabel="Mean code value")
    axes[0, 1].set(title="Noise versus signal", xlabel="Mean code value", ylabel="Spatial noise σ")
    axes[1, 0].set(title="SNR versus signal", xlabel="Mean code value", ylabel="SNR (dB)")
    axes[1, 1].set(title="Repeat consistency", xlabel="Gray patch ID", ylabel="Between-repeat σ")
    axes[0, 0].legend(fontsize=8, ncol=2)
    for ax in axes.flat:
        ax.grid(True, alpha=0.25)
    fig.suptitle("Grayscale chart analysis")
    fig.savefig(output, dpi=180)
    plt.close(fig)


def plot_color(summary: pd.DataFrame, output: Path) -> None:
    lux_summary = summary.groupby("lux", as_index=False).agg(
        mean_chroma_retention=("chroma_retention_vs_max_lux", "mean"),
        mean_spatial_noise=("spatial_noise_std_mean", "mean"),
        mean_snr_db=("snr_db_mean", "mean"),
    )
    fig, axes = plt.subplots(2, 2, figsize=(12, 9), constrained_layout=True)
    axes[0, 0].plot(lux_summary["lux"], lux_summary["mean_chroma_retention"], marker="o")
    axes[0, 1].plot(lux_summary["lux"], lux_summary["mean_spatial_noise"], marker="o")
    axes[1, 0].plot(lux_summary["lux"], lux_summary["mean_snr_db"], marker="o")
    if "delta_E76" in summary:
        de = summary.groupby("lux", as_index=False)["delta_E76"].mean()
        axes[1, 1].plot(de["lux"], de["delta_E76"], marker="o")
        axes[1, 1].set(title="Mean ΔE76", ylabel="ΔE76")
    else:
        axes[1, 1].text(0.5, 0.5, "Provide --color-reference\nfor ΔE76", ha="center", va="center")
        axes[1, 1].set_axis_off()
    for ax, title, ylabel in [
        (axes[0, 0], "Chroma retention", "Ratio to maximum-lux chroma"),
        (axes[0, 1], "Colour-patch noise", "Mean spatial noise σ"),
        (axes[1, 0], "Colour-patch SNR", "Mean SNR (dB)"),
    ]:
        ax.set_title(title); ax.set_xlabel("Illuminance (lux, log scale)"); ax.set_ylabel(ylabel); ax.set_xscale("log")
    if axes[1, 1].axison:
        axes[1, 1].set_xlabel("Illuminance (lux, log scale)"); axes[1, 1].set_xscale("log")
    for ax in axes.flat:
        if ax.axison: ax.grid(True, alpha=0.25)
    fig.suptitle("Colour chart analysis")
    fig.savefig(output, dpi=180)
    plt.close(fig)


def style_workbook(path: Path) -> None:
    wb = load_workbook(path)
    for ws in wb.worksheets:
        ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions; ws.sheet_view.showGridLines = False
        for cell in ws[1]:
            cell.fill = PatternFill("solid", fgColor="17365D")
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.row_dimensions[1].height = 36
        for col in range(1, ws.max_column + 1):
            values = [str(ws.cell(row, col).value or "") for row in range(1, min(ws.max_row, 100) + 1)]
            ws.column_dimensions[get_column_letter(col)].width = min(34, max(10, max(map(len, values), default=0) + 2))
        for row in range(2, ws.max_row + 1, 2):
            for cell in ws[row]: cell.fill = PatternFill("solid", fgColor="EAF2F8")
    wb.save(path)


def main() -> None:
    args = parse_args()
    if args.write_default_config:
        write_config(args.config); return
    config = load_config(args.config)
    root, output = args.input.resolve(), args.output.resolve()
    gray_groups = collect_groups(root / "grayscale")
    color_groups = collect_groups(root / "colorchecker")
    if not gray_groups and not color_groups:
        raise SystemExit("No lux folders found below grayscale/ or colorchecker/")
    output.mkdir(parents=True, exist_ok=True)
    diagnostic_dir = output / "diagnostic_patch_overlays"; diagnostic_dir.mkdir(exist_ok=True)
    gray_details = color_details = gray_summary = color_summary = pd.DataFrame()
    if gray_groups:
        gray_details = analyze_chart("grayscale", gray_groups, config["grayscale"],
                                     float(config["patch_inset_fraction"]), bool(config.get("align_repeats", True)), diagnostic_dir)
        gray_ref = load_reference(args.gray_reference, ["patch_id"]) if args.gray_reference else None
        gray_summary = summarize_gray(gray_details, gray_ref)
        gray_details.to_csv(output / "grayscale_per_image_patch_metrics.csv", index=False, float_format="%.6f")
        gray_summary.to_csv(output / "grayscale_summary.csv", index=False, float_format="%.6f")
        plot_gray(gray_summary, output / "grayscale_analysis.png")
    if color_groups:
        color_details = analyze_chart("colorchecker", color_groups, config["colorchecker"],
                                      float(config["patch_inset_fraction"]), bool(config.get("align_repeats", True)), diagnostic_dir)
        color_ref = load_reference(args.color_reference, ["patch_id", "L_ref", "a_ref", "b_ref"]) if args.color_reference else None
        color_summary = summarize_color(color_details, color_ref)
        color_details.to_csv(output / "color_per_image_patch_metrics.csv", index=False, float_format="%.6f")
        color_summary.to_csv(output / "color_summary.csv", index=False, float_format="%.6f")
        plot_color(color_summary, output / "color_analysis.png")
    method = pd.DataFrame([
        ["Method status", "Separate-chart, ISO 19093-inspired analysis; not a claim of conformity"],
        ["Input", str(root)], ["Configuration", str(args.config.resolve())],
        ["Patch sampling", f"Interior ROI with {100*float(config['patch_inset_fraction']):g}% inset on every edge"],
        ["Repeat alignment", "Translation ECC after chart perspective correction" if config.get("align_repeats", True) else "Disabled"],
        ["Spatial noise", "Sample standard deviation of grayscale pixels inside a flat patch"],
        ["High-pass noise", "Sample standard deviation after subtracting Gaussian-blurred patch"],
        ["SNR", "20 log10(mean grayscale code value / spatial noise standard deviation)"],
        ["OECF x-axis", "log10(lux × relative patch luminance); proxy only unless patch luminance is calibrated"],
        ["Colour conversion", "Output-referred JPEG/TIFF interpreted as sRGB, converted to CIELAB"],
        ["Colour error", "ΔE76; calculated only when calibrated reference Lab values are supplied"],
        ["Chroma retention", "Patch chroma divided by that patch's chroma at the highest measured lux"],
    ], columns=["Setting", "Value"])
    excel_path = output / "chart_analysis.xlsx"
    with pd.ExcelWriter(excel_path, engine="openpyxl") as writer:
        if not gray_summary.empty: gray_summary.to_excel(writer, sheet_name="Gray Summary", index=False)
        if not gray_details.empty: gray_details.to_excel(writer, sheet_name="Gray Per Image", index=False)
        if not color_summary.empty: color_summary.to_excel(writer, sheet_name="Color Summary", index=False)
        if not color_details.empty: color_details.to_excel(writer, sheet_name="Color Per Image", index=False)
        method.to_excel(writer, sheet_name="Method", index=False)
    style_workbook(excel_path)
    print(f"Grayscale images: {sum(len(files) for _, _, files in gray_groups)}")
    print(f"Colour-chart images: {sum(len(files) for _, _, files in color_groups)}")
    print(f"Results: {output}")


if __name__ == "__main__":
    main()
